from pathlib import Path
from datetime import date
import os

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "docs"
ASSET_DIR = OUT_DIR / "week6_report_assets"
OUTPUT = Path(os.environ.get(
    "WEEK6_REPORT_OUTPUT",
    str(OUT_DIR / "6주차_HTTPS_트래픽_분석_기능_구현_보고서.docx"),
))

ASSET_DIR.mkdir(parents=True, exist_ok=True)

FONT_KR = "Malgun Gothic"
FONT_MONO = "Consolas"
NAVY = "17365D"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
MUTED = "666666"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
CALLOUT = "F4F6F9"
GREEN = "2F6B4F"
RED = "9B1C1C"
GOLD = "7A5A00"
WHITE = "FFFFFF"


def rgb(hex_value):
    return RGBColor.from_string(hex_value)


def pil_color(hex_value):
    if isinstance(hex_value, str) and len(hex_value) == 6 and not hex_value.startswith("#"):
        return f"#{hex_value}"
    return hex_value


def set_run_font(run, name=FONT_KR, size=11, color="000000", bold=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT_KR)
    run.font.size = Pt(size)
    run.font.color.rgb = rgb(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa, indent_dxa=120):
    total = sum(widths_dxa)
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")

    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            width = widths_dxa[min(idx, len(widths_dxa) - 1)]
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_table_borders(table, color="C7CED8", size="4"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_repeat_keep(paragraph, keep_next=False, keep_lines=False):
    paragraph.paragraph_format.keep_with_next = keep_next
    paragraph.paragraph_format.keep_together = keep_lines


def add_field(paragraph, instruction):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])
    set_run_font(run, size=9, color=MUTED)


def add_numbering_definition(doc, num_fmt, text, left=720, hanging=360):
    numbering = doc.part.numbering_part.element
    abstract_ids = [int(x.get(qn("w:abstractNumId"))) for x in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(x.get(qn("w:numId"))) for x in numbering.findall(qn("w:num"))]
    abstract_id = max(abstract_ids, default=-1) + 1
    num_id = max(num_ids, default=0) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    fmt = OxmlElement("w:numFmt")
    fmt.set(qn("w:val"), num_fmt)
    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), text)
    suff = OxmlElement("w:suff")
    suff.set(qn("w:val"), "tab")
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), str(left))
    tabs.append(tab)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), str(left))
    ind.set(qn("w:hanging"), str(hanging))
    p_pr.extend([tabs, ind])
    r_pr = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), FONT_KR)
    fonts.set(qn("w:hAnsi"), FONT_KR)
    fonts.set(qn("w:eastAsia"), FONT_KR)
    r_pr.append(fonts)
    lvl.extend([start, fmt, lvl_text, suff, p_pr, r_pr])
    abstract.append(lvl)
    numbering.append(abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)
    return num_id


def apply_num(paragraph, num_id):
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = p_pr.find(qn("w:numPr"))
    if num_pr is None:
        num_pr = OxmlElement("w:numPr")
        p_pr.append(num_pr)
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_id_el = OxmlElement("w:numId")
    num_id_el.set(qn("w:val"), str(num_id))
    num_pr.extend([ilvl, num_id_el])


def add_bullet(doc, text, bullet_num_id, bold_prefix=None):
    p = doc.add_paragraph()
    apply_num(p, bullet_num_id)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.167
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        set_run_font(r1, bold=True)
        r2 = p.add_run(text[len(bold_prefix):])
        set_run_font(r2)
    else:
        set_run_font(p.add_run(text))
    return p


def add_numbered(doc, text, num_id):
    p = doc.add_paragraph()
    apply_num(p, num_id)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.167
    set_run_font(p.add_run(text))
    return p


def add_body(doc, text, bold_prefix=None, keep=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.10
    p.paragraph_format.keep_together = keep
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        set_run_font(r1, bold=True)
        r2 = p.add_run(text[len(bold_prefix):])
        set_run_font(r2)
    else:
        set_run_font(p.add_run(text))
    return p


def add_callout(doc, label, text, fill=CALLOUT, accent=BLUE):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [9360])
    set_table_borders(table, color=accent, size="6")
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.10
    r = p.add_run(f"{label}  ")
    set_run_font(r, size=10.5, color=accent, bold=True)
    set_run_font(p.add_run(text), size=10.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_code_block(doc, lines):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [9360])
    set_table_borders(table, color="D4D8DE", size="4")
    cell = table.cell(0, 0)
    set_cell_shading(cell, "F7F8FA")
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    for idx, line in enumerate(lines):
        if idx:
            p.add_run().add_break()
        run = p.add_run(line)
        set_run_font(run, name=FONT_MONO, size=8.2, color="263238")
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(8)
    set_run_font(p.add_run(text), size=9, color=MUTED, italic=True)
    return p


def style_table(table, header=True, body_size=9.4):
    set_table_borders(table)
    for r_idx, row in enumerate(table.rows):
        for cell in row.cells:
            if header and r_idx == 0:
                set_cell_shading(cell, LIGHT_GRAY)
            for p in cell.paragraphs:
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.08
                for run in p.runs:
                    set_run_font(run, size=body_size, bold=(header and r_idx == 0), color=NAVY if header and r_idx == 0 else "000000")
    if header:
        repeat_table_header(table.rows[0])


def add_matrix(doc, headers, rows, widths, font_size=9.2):
    table = doc.add_table(rows=1, cols=len(headers))
    for idx, header in enumerate(headers):
        table.rows[0].cells[idx].text = header
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            cells[idx].text = value
    set_table_geometry(table, widths)
    style_table(table, body_size=font_size)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def load_font(size, bold=False):
    paths = [
        Path("C:/Windows/Fonts/malgunbd.ttf" if bold else "C:/Windows/Fonts/malgun.ttf"),
        Path("C:/Windows/Fonts/arialbd.ttf" if bold else "C:/Windows/Fonts/arial.ttf"),
    ]
    for path in paths:
        if path.exists():
            return ImageFont.truetype(str(path), size)
    return ImageFont.load_default()


def rounded_box(draw, rect, fill, outline, title, subtitle=None, title_color="17365D"):
    draw.rounded_rectangle(rect, radius=22, fill=pil_color(fill), outline=pil_color(outline), width=3)
    x1, y1, x2, y2 = rect
    title_font = load_font(31, bold=True)
    body_font = load_font(22)
    title_bbox = draw.textbbox((0, 0), title, font=title_font)
    title_w = title_bbox[2] - title_bbox[0]
    title_y = y1 + 25 if subtitle else y1 + (y2 - y1 - (title_bbox[3] - title_bbox[1])) // 2
    draw.text(((x1 + x2 - title_w) / 2, title_y), title, font=title_font, fill=pil_color(title_color))
    if subtitle:
        sub_bbox = draw.textbbox((0, 0), subtitle, font=body_font)
        sub_w = sub_bbox[2] - sub_bbox[0]
        draw.text(((x1 + x2 - sub_w) / 2, title_y + 48), subtitle, font=body_font, fill="#4B5563")


def arrow(draw, start, end, color="2E74B5", width=6):
    draw.line([start, end], fill=pil_color(color), width=width)
    ex, ey = end
    sx, sy = start
    if abs(ex - sx) >= abs(ey - sy):
        direction = 1 if ex > sx else -1
        pts = [(ex, ey), (ex - direction * 18, ey - 12), (ex - direction * 18, ey + 12)]
    else:
        direction = 1 if ey > sy else -1
        pts = [(ex, ey), (ex - 12, ey - direction * 18), (ex + 12, ey - direction * 18)]
    draw.polygon(pts, fill=pil_color(color))


def create_architecture_diagram(path):
    img = Image.new("RGB", (1800, 620), "#FFFFFF")
    draw = ImageDraw.Draw(img)
    boxes = [
        ((40, 205, 330, 405), "Browser", "HTTPS 요청 / 파일 업로드", "E8EEF5"),
        ((410, 130, 760, 480), "Local DLP Proxy", "CONNECT 정책 · TLS 종단", "F4F6F9"),
        ((850, 205, 1190, 405), "HTTP 분석 계층", "URL · Header · Body", "FFF8E8"),
        ((1280, 205, 1760, 405), "AI Service", "ChatGPT / Gemini / Claude", "EAF4EE"),
    ]
    for rect, title, subtitle, fill in boxes:
        rounded_box(draw, rect, fill, "7F8FA6", title, subtitle)
    arrow(draw, (330, 305), (410, 305))
    arrow(draw, (760, 305), (850, 305))
    arrow(draw, (1190, 305), (1280, 305))
    small = load_font(21, bold=True)
    draw.text((345, 250), "TLS A", font=small, fill=pil_color(BLUE))
    draw.text((1208, 250), "TLS B", font=small, fill=pil_color(GREEN))
    draw.text((905, 450), "복호화 평문을 요청/응답 단위로 검사 후 재전송", font=load_font(23), fill=pil_color(NAVY))
    img.save(path)


def create_connect_diagram(path):
    img = Image.new("RGB", (1750, 1040), "#FFFFFF")
    draw = ImageDraw.Draw(img)
    title_font = load_font(35, bold=True)
    body_font = load_font(23)
    draw.text((65, 35), "HTTPS CONNECT 요청 처리 흐름", font=title_font, fill=pil_color(NAVY))
    nodes = [
        ((610, 110, 1140, 225), "1. CONNECT host:443 수신", "Request Buffer에서 헤더 경계 확인"),
        ((610, 285, 1140, 400), "2. 대상·프로세스 식별", "Host/Port 해석, 자기 자신 연결 방지"),
        ((610, 460, 1140, 575), "3. TLS 정책 결정", "MITM / BYPASS / BLOCK / AUDIT / IGNORE"),
        ((90, 680, 530, 830), "Raw Tunnel", "비 AI 사이트는 암호화 바이트만 전달"),
        ((655, 680, 1095, 830), "TLS MITM", "200 응답 후 이중 TLS 세션 생성"),
        ((1220, 680, 1660, 830), "Block", "정책 사유와 함께 연결 거부"),
    ]
    for idx, (rect, title, sub) in enumerate(nodes):
        fill = "F4F6F9"
        outline = "7F8FA6"
        if idx == 3:
            fill, outline = "EDF2F7", "607D8B"
        elif idx == 4:
            fill, outline = "E8EEF5", BLUE
        elif idx == 5:
            fill, outline = "FCEBEC", RED
        rounded_box(draw, rect, fill, outline, title, sub)
    arrow(draw, (875, 225), (875, 285))
    arrow(draw, (875, 400), (875, 460))
    arrow(draw, (740, 575), (310, 680))
    arrow(draw, (875, 575), (875, 680))
    arrow(draw, (1010, 575), (1440, 680))
    draw.text((700, 605), "정책 결과에 따른 분기", font=body_font, fill=pil_color(MUTED))
    draw.text((145, 890), "기본값: IGNORE (비 AI 트래픽 무검사 통과)", font=body_font, fill="#37474F")
    draw.text((650, 890), "AI 브라우저 트래픽: MITM + 파일 업로드 분석", font=body_font, fill=pil_color(NAVY))
    img.save(path)


def configure_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT_KR
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT_KR)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT_KR)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_KR)
    normal.font.size = Pt(11)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ):
        style = styles[name]
        style.font.name = FONT_KR
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT_KR)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT_KR)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_KR)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = rgb(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True

    if "Code Block" not in styles:
        code_style = styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
    else:
        code_style = styles["Code Block"]
    code_style.font.name = FONT_MONO
    code_style._element.rPr.rFonts.set(qn("w:ascii"), FONT_MONO)
    code_style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT_MONO)
    code_style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_KR)
    code_style.font.size = Pt(8.2)


def configure_section(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    header = section.header
    hp = header.paragraphs[0]
    hp.paragraph_format.space_after = Pt(0)
    hp.paragraph_format.tab_stops.add_tab_stop(Inches(6.5))
    set_run_font(hp.add_run("Local DLP HTTPS Proxy PoC"), size=8.5, color=MUTED, bold=True)
    hp.add_run("\t")
    set_run_font(hp.add_run("6주차 기술 진행 보고"), size=8.5, color=MUTED)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fp.paragraph_format.space_before = Pt(0)
    set_run_font(fp.add_run("Page "), size=9, color=MUTED)
    add_field(fp, "PAGE")
    set_run_font(fp.add_run(" / "), size=9, color=MUTED)
    add_field(fp, "NUMPAGES")
    return section


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    set_repeat_keep(p, keep_next=True, keep_lines=True)
    return p


def add_title_page_content(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(3)
    set_run_font(p.add_run("WEEK 06 · TECHNICAL PROGRESS REPORT"), size=9.5, color=BLUE, bold=True)

    title = doc.add_paragraph()
    title.paragraph_format.space_before = Pt(4)
    title.paragraph_format.space_after = Pt(4)
    title.paragraph_format.keep_with_next = True
    set_run_font(title.add_run("6주차 HTTPS 트래픽 분석\n기능 구현 보고서"), size=25, color=NAVY, bold=True)

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(18)
    set_run_font(
        subtitle.add_run("TLS 복호화 이후 평문 분석 구조와 ChatGPT 파일 업로드 검증 결과"),
        size=13,
        color=MUTED,
    )

    metadata = [
        ("보고 대상", "선임 개발자"),
        ("프로젝트", "Local DLP HTTPS Proxy PoC"),
        ("작성일", "2026년 7월 15일"),
        ("보고 범위", "HTTPS 트래픽 분석 기능 1~5번"),
        ("현재 상태", "ChatGPT HTTP/2 파일 업로드 복호화·연결·분석·전송 검증 완료"),
    ]
    for label, value in metadata:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.0
        set_run_font(p.add_run(f"{label}: "), size=10.5, color=NAVY, bold=True)
        set_run_font(p.add_run(value), size=10.5)

    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    add_callout(
        doc,
        "이번 주 핵심 성과",
        "브라우저의 CONNECT 요청을 AI 사이트에 한정해 MITM 처리하고, ALPN으로 HTTP/1.1과 HTTP/2를 분기한 뒤 ChatGPT 메타데이터 요청과 raw PUT 스트림을 연결하여 한글 파일명, DOCX 형식, SHA-256, 추출 바이트 수, 전송 결과를 하나의 업로드 ID로 확인했다.",
        fill=LIGHT_BLUE,
        accent=BLUE,
    )


def build_report():
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    architecture_path = ASSET_DIR / "https_mitm_architecture.png"
    connect_path = ASSET_DIR / "connect_processing_flow.png"
    create_architecture_diagram(architecture_path)
    create_connect_diagram(connect_path)

    doc = Document()
    configure_styles(doc)
    configure_section(doc)
    bullet_num = add_numbering_definition(doc, "bullet", "•", left=720, hanging=360)
    decimal_num = add_numbering_definition(doc, "decimal", "%1.", left=720, hanging=360)

    core = doc.core_properties
    core.title = "6주차 HTTPS 트래픽 분석 기능 구현 보고서"
    core.subject = "Local DLP HTTPS Proxy PoC 주간 기술 진행 보고"
    core.author = "Local DLP Proxy 개발 담당"
    core.keywords = "HTTPS, TLS MITM, HTTP/2, CONNECT, DLP, 파일 업로드"

    add_title_page_content(doc)
    add_heading(doc, "요약", level=1)
    add_body(
        doc,
        "이번 주에는 기존 HTTP Local Proxy를 HTTPS 트래픽 분석 구조로 확장했다. 구현의 중심은 CONNECT 요청을 단순 터널로만 전달하지 않고, 정책에 따라 Proxy가 TLS를 종단하여 평문 HTTP 메시지를 분석한 뒤 별도의 TLS 세션으로 원격 서버에 재전송하는 것이다. 비 AI 웹사이트와 데스크톱 애플리케이션은 DEFAULT IGNORE 정책으로 암호화 상태 그대로 통과시키고, Chrome·Edge·Firefox·Brave가 ChatGPT·Gemini·Claude로 연결할 때만 MITM과 파일 업로드 검사를 수행하도록 범위를 축소했다."
    )
    add_body(
        doc,
        "실환경 검증에서는 ChatGPT가 HTTP/2를 선택하는 것을 ALPN 로그로 확인했다. 파일 등록 메타데이터 요청에서 원본 파일명을 얻고, 별도 oaiusercontent.com 호스트로 전송되는 raw PUT 본문과 크기·경로를 연결했다. DOCX 업로드는 ZIP 구조, 내부 엔트리, 텍스트 추출 가능 여부, SHA-256을 검사한 후 안전한 파일을 upstream_status=201로 전달했다."
    )

    add_heading(doc, "1. 이번 주 진행 범위 및 결과", level=1)
    add_matrix(
        doc,
        ["No.", "구현 항목", "상태", "이번 주 확인 결과"],
        [
            ("1", "TLS 복호화 이후 Proxy 내부 데이터 흐름", "완료", "TLS read → HTTP 경계 판별 → 파싱 → 파일 검사 → TLS write 흐름 구성"),
            ("2", "HTTPS CONNECT 요청 처리 구조", "완료", "대상 해석·정책 결정·200 응답·MITM/Raw/Block 분기 구현"),
            ("3", "Client-Proxy / Proxy-Server 이중 TLS", "완료", "동적 인증서, SNI, upstream 인증서 검증, ALPN 전달 구현"),
            ("4", "복호화된 요청/응답 평문 분석", "완료", "HTTP/1.1 파서와 HTTP/2 스트림 분석기를 프로토콜에 따라 분기"),
            ("5", "URL·Header·Body 추출", "완료", "method/path/host/content-type/content-length/header 배열/full body 추출"),
        ],
        [600, 2500, 900, 5360],
        font_size=9.1,
    )

    doc.add_page_break()
    add_heading(doc, "2. 전체 구조", level=1)
    add_body(
        doc,
        "클라이언트가 HTTPS 사이트에 접속하면 Proxy에는 먼저 암호화되지 않은 CONNECT host:port 요청이 전달된다. Proxy는 이 시점에 프로세스명과 목적지 호스트를 기준으로 TLS 가로채기 정책을 결정한다. MITM 대상이라면 CONNECT 200 응답 이후 브라우저와의 TLS 세션 A, 원격 AI 서버와의 TLS 세션 B를 각각 구성한다. 두 TLS 세션 사이에서만 HTTP 평문이 존재하며, 이 지점이 URL·Header·Body 추출과 DLP 파일 분석의 실행 위치다."
    )
    doc.add_picture(str(architecture_path), width=Inches(6.45))
    add_caption(doc, "그림 1. TLS 복호화 이후 Proxy 내부 데이터 흐름")
    add_callout(
        doc,
        "설계 원칙",
        "모든 HTTPS를 복호화하지 않는다. AI 사이트의 브라우저 파일 업로드만 검사하고, 일반 웹사이트·인증 서비스·백그라운드 트래픽은 암호화된 raw tunnel로 통과시킨다.",
        fill="EEF4F8",
        accent=DARK_BLUE,
    )

    add_heading(doc, "3. 항목별 구현 내용", level=1)
    add_heading(doc, "3.1 TLS 복호화 이후 Proxy 내부 데이터 흐름", level=2)
    add_body(
        doc,
        "복호화 이후의 데이터 경로는 단순한 recv/send가 아니라 프로토콜별 메시지 경계와 정책 결과를 보존해야 한다. HTTP/1.1에서는 request_buffer와 response_buffer가 완전한 메시지 길이를 판단하고, HTTP/2에서는 stream_id 단위 상태 객체가 HEADERS와 DATA 프레임을 결합한다. 공통 분석 결과는 ALLOW 또는 BLOCK으로 정규화한 뒤 upstream 또는 client 방향으로 다시 TLS 암호화해 전달한다."
    )
    for step in [
        "Client TLS에서 SSL_read()로 복호화된 바이트를 수신한다.",
        "협상된 ALPN이 h2이면 http2_engine_relay_loop(), 아니면 HTTP/1.1 요청/응답 루프로 진입한다.",
        "요청 경계를 확인한 뒤 method, path, host, headers, body를 http_request_t에 저장한다.",
        "파일 업로드 요청이면 메타데이터와 raw PUT을 연결하고 파일 형식·해시·내부 구조를 검사한다.",
        "ALLOW이면 upstream SSL_write(), BLOCK이면 해당 스트림 또는 요청에 차단 응답을 보낸다.",
        "응답도 동일하게 경계를 판별하고 필요한 분석을 수행한 후 client TLS로 재암호화한다.",
    ]:
        add_numbered(doc, step, decimal_num)
    add_code_block(
        doc,
        [
            "Client SSL_read → Request/HTTP2 Stream Buffer → HTTP Parser",
            "→ DLP/File Analyzer → Upstream SSL_write",
            "Upstream SSL_read → Response Parser → Client SSL_write",
        ],
    )
    add_body(
        doc,
        "주요 구현 파일은 tls_mitm_engine.c, http2_engine.c, request_buffer.c, response_buffer.c, dlp_engine.c, file_analyzer.c다. TLS 계층은 암·복호화와 프로토콜 선택을 담당하고, HTTP 계층은 메시지 구조를 복원하며, DLP 계층은 실제 파일 여부와 허용·차단 결과를 결정한다."
    )

    add_heading(doc, "3.2 HTTPS CONNECT 요청 처리 구조", level=2)
    add_body(
        doc,
        "relay_proxy.c의 CONNECT 처리 경로는 요청 대상 해석, 정책 결정, upstream TCP 연결, CONNECT 200 응답, 실행 모드 분기의 순서로 구성된다. CONNECT 자체에는 실제 HTTPS URL 경로나 파일 본문이 없으므로, 이 단계의 핵심은 host:port와 클라이언트 프로세스를 정확히 식별하고 이후 트래픽 처리 모드를 안전하게 선택하는 것이다."
    )
    doc.add_picture(str(connect_path), width=Inches(6.35))
    add_caption(doc, "그림 2. CONNECT 요청 처리 및 정책 분기")
    add_matrix(
        doc,
        ["정책", "처리 방식", "적용 목적"],
        [
            ("MITM", "CONNECT 200 이후 이중 TLS 구성", "AI 사이트 파일 업로드 복호화·검사"),
            ("BYPASS", "암호화 바이트 단순 전달", "명시적 예외 사이트"),
            ("IGNORE", "로그를 억제한 raw tunnel", "비 AI 사이트와 알 수 없는 목적지 기본값"),
            ("AUDIT", "내용은 복호화하지 않고 연결 정보만 기록", "단계적 정책 도입"),
            ("BLOCK", "CONNECT 단계에서 연결 거부", "명시적 금지 대상"),
        ],
        [1100, 3300, 4960],
        font_size=9.2,
    )
    add_callout(
        doc,
        "안전 기본값",
        "현재 tls_intercept_policy.txt의 DEFAULT는 IGNORE다. 정책에 없는 HTTPS는 내용 분석이나 차단 없이 기존 사용자 경험을 유지한다.",
        fill="F4F6F9",
        accent=DARK_BLUE,
    )

    add_heading(doc, "3.3 Client-Proxy TLS 세션과 Proxy-Server TLS 세션", level=2)
    add_body(
        doc,
        "TLS MITM은 하나의 TLS 연결을 해제하는 기능이 아니라 두 개의 독립된 TLS 연결을 동시에 유지하는 구조다. Client-Proxy 구간에서 Proxy는 서버 역할을 수행하며 SSL_accept()로 브라우저 핸드셰이크를 수락한다. SNI 콜백은 접속 호스트에 맞는 동적 leaf 인증서를 선택하고, 해당 인증서는 Local DLP MITM Root CA로 서명된다."
    )
    add_body(
        doc,
        "Proxy-Server 구간에서 Proxy는 클라이언트 역할로 SSL_connect()를 수행한다. 브라우저에서 추출한 SNI를 upstream에 전달하고, Windows 신뢰 저장소를 OpenSSL 인증서 저장소로 로드하여 실제 원격 서버 인증서를 검증한다. 운영 연결에서는 LOCAL_DLP_ALLOW_INSECURE_UPSTREAM을 제거하여 인증서 검증을 우회하지 않도록 했다."
    )
    add_matrix(
        doc,
        ["구분", "Client ↔ Proxy (TLS A)", "Proxy ↔ AI Server (TLS B)"],
        [
            ("Proxy 역할", "TLS Server", "TLS Client"),
            ("핸드셰이크 함수", "SSL_accept()", "SSL_connect()"),
            ("인증서", "동적 leaf + Local Root CA", "원격 서비스 실제 인증서"),
            ("SNI", "브라우저 SNI 수신", "동일 SNI를 upstream에 설정"),
            ("ALPN", "h2 또는 http/1.1 선택", "클라이언트와 동일 프로토콜 요청"),
            ("검증 실패 시", "브라우저가 연결 거부", "Proxy가 upstream 연결 중단"),
        ],
        [1300, 4030, 4030],
        font_size=9.0,
    )
    add_code_block(
        doc,
        [
            "[INFO] AI TLS READY session=8 host=chatgpt.com protocol=h2",
            "[INFO] AI TLS READY session=41 host=files.openai.com protocol=h2",
            "[INFO] AI TLS READY session=44 host=sdmntprseasia.oaiusercontent.com protocol=h2",
        ],
    )
    add_body(
        doc,
        "위 로그는 Client-Proxy와 Proxy-Server 핸드셰이크가 모두 완료되고 양쪽이 HTTP/2를 협상한 뒤에만 기록된다. 따라서 AI TLS READY는 단순 TCP 접속 성공이 아니라 실제 복호화 분석 루프 진입 가능 상태를 의미한다."
    )

    add_heading(doc, "3.4 복호화된 HTTPS 요청/응답의 HTTP 평문 분석", level=2)
    add_body(
        doc,
        "TLS 핸드셰이크가 끝나면 SSL_read() 결과는 HTTP 평문이 된다. HTTP/1.1 경로에서는 tls_mitm_process_one_http_request_response()가 한 번의 요청·응답 교환을 처리한다. request_buffer_get_complete_request_length()와 response_buffer_get_complete_response_length()가 Content-Length, chunked 여부, 헤더 종료 위치를 이용해 완전한 메시지 경계를 확인하고, 이후 parse_http_request()와 parse_http_response()를 호출한다."
    )
    add_body(
        doc,
        "HTTP/2 경로에서는 http2_engine_relay_loop()가 프레임 헤더를 파싱하고 HPACK으로 HEADERS 블록을 해제한다. stream_id마다 method, path, authority, content-type, content-length, body 누적량과 분석 상태를 보관하기 때문에 여러 요청이 하나의 TLS 연결에서 동시에 진행되어도 업로드 단위를 구분할 수 있다."
    )
    add_matrix(
        doc,
        ["분석 단계", "HTTP/1.1", "HTTP/2"],
        [
            ("메시지 단위", "요청/응답 순서", "stream_id"),
            ("Header 복원", "CRLF 라인 파싱", "HPACK decode"),
            ("Body 수집", "Content-Length 또는 chunked", "DATA 프레임 누적"),
            ("다중 요청", "Keep-Alive 순차 처리", "Multiplexing 동시 처리"),
            ("차단 단위", "HTTP 응답/연결", "RST_STREAM 또는 스트림 차단"),
        ],
        [1500, 3930, 3930],
        font_size=9.2,
    )
    add_callout(
        doc,
        "현재 분석 정책",
        "일반 채팅 JSON, 텔레메트리, 응답 본문은 파일 분석 대상에서 제외한다. dlp_request_should_inspect()가 실제 파일 업로드로 판단한 요청만 DLP와 파일 분석기로 전달한다.",
        fill="F4F6F9",
        accent=BLUE,
    )

    add_heading(doc, "3.5 URL, Header, Body 데이터 추출 방식", level=2)
    add_body(
        doc,
        "HTTP/1.1의 parse_http_request()는 시작 줄에서 method, path, version을 추출하고, 이후 각 헤더 라인을 배열에 보관한다. Host, Content-Length, Content-Type은 DLP에서 자주 사용하므로 별도 필드에도 저장한다. body_data는 완전한 요청 버퍼를 가리켜 대용량 파일 분석에 사용하고, body는 로그와 간단한 판정을 위한 제한된 미리보기다."
    )
    add_matrix(
        doc,
        ["데이터", "HTTP/1.1 추출 위치", "HTTP/2 추출 위치", "활용"],
        [
            ("URL/Path", "Request Line 두 번째 토큰", ":path pseudo-header", "업로드 API·raw 경로 식별"),
            ("Host", "Host 헤더", ":authority pseudo-header", "AI/업로드 호스트 정책"),
            ("Method", "Request Line 첫 토큰", ":method pseudo-header", "POST/PUT 업로드 판정"),
            ("Content-Type", "Content-Type 헤더", "content-type header", "DOCX/PDF/ZIP 형식 판단"),
            ("Content-Length", "정수 헤더 값", "content-length header", "메타데이터와 raw PUT 연결"),
            ("Body", "헤더 종료 뒤 전체 바이트", "DATA 프레임 누적", "해시·압축 해제·파일 검사"),
        ],
        [1150, 2500, 2500, 3210],
        font_size=8.8,
    )
    add_body(
        doc,
        "ChatGPT 업로드는 원본 파일명이 포함된 메타데이터 요청과 실제 파일 바이트를 전송하는 raw PUT 요청이 서로 다른 호스트·스트림에서 발생한다. upload_tracker_record_metadata_request()가 file_name, mime_type, file_size를 저장하고, upload_tracker_match_raw_put()가 업로드 호스트·경로·크기를 기준으로 실제 스트림에 원본 파일명을 연결한다."
    )
    add_code_block(
        doc,
        [
            "UPLOAD PREPARED id=1 file=\"5주차_TLS_MITM_구현사항_반영_정리.docx\"",
            "  declared_bytes=46494 type=application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            "UPLOAD INSPECTED id=1 session=44 stream=3 bytes=46494 format=DOCX entries=18",
            "  extracted_text_bytes=14163 sha256=85b0...421 action=ALLOW",
            "UPLOAD FORWARDED id=1 session=44 stream=3 upstream_status=201",
        ],
    )

    doc.add_page_break()
    add_heading(doc, "4. 구현 중 발생한 문제와 수정 내용", level=1)
    add_body(
        doc,
        "초기 구현 과정에서는 기능 자체보다 연결 안정성, 정책 범위, 로그 가독성에서 문제가 크게 나타났다. 아래 표는 실제 증상과 수정 지점을 요약한 것이다."
    )
    add_matrix(
        doc,
        ["문제/현상", "발견 로그·원인", "수정 코드", "결과"],
        [
            ("브라우저가 ChatGPT 화면을 불완전하게 표시", "프록시 재시작 시 기존 TLS/H2 연결이 종료되고 SSL_accept·SSL_read가 ERROR로 누적", "tls_mitm_engine.c, http2_engine.c에서 정상 EOF·10053·10054를 재연결 종료로 분류", "불필요한 ERROR 연쇄 제거, 실제 오류에 상세 원인 추가"),
            ("모든 사이트·애플리케이션 로그가 과다 생성", "기본 MITM 범위가 넓고 일반 요청까지 INFO 출력", "tls_intercept_policy.txt를 브라우저+AI 호스트 규칙으로 축소, DEFAULT IGNORE; logger.c INFO 필터 추가", "일반 사이트는 무검사 통과, 파일 업로드 중심 로그 확보"),
            ("ChatGPT 업로드 원본 파일명 확인 불가", "메타데이터 요청과 oaiusercontent raw PUT이 별도 스트림", "upload_tracker.c 및 http2_engine.c에서 upload_id, 크기, 경로로 상관관계 구성", "원본 파일명과 실제 바이트 분석 결과를 하나의 로그로 연결"),
            ("한글 파일명이 콘솔에서 깨짐", "로그 파일은 UTF-8 정상이나 printf 콘솔이 CP949로 해석", "upload_tracker.c의 JSON Unicode→UTF-8 복원 유지, logger.c에서 UTF-8→UTF-16 후 WriteConsoleW 사용", "한글 파일명 데이터 보존 및 콘솔 표시 개선"),
            ("upstream 인증서 검증 실패 위험", "OpenSSL이 Windows Root Store를 자동 사용하지 않음", "tls_mitm_engine.c에서 CurrentUser/LocalMachine ROOT를 OpenSSL X509_STORE에 로드", "운영 연결에서 insecure 우회 없이 원격 인증서 검증"),
            ("포트 중복 실행", "bind() failed: port 8000 is already in use", "exclusive listener와 PID/포트 진단 로그", "중복 프로세스를 즉시 식별"),
        ],
        [1800, 2670, 2870, 2020],
        font_size=8.3,
    )

    add_heading(doc, "4.1 대표 장애 로그와 해석", level=2)
    add_code_block(
        doc,
        [
            "[ERROR] HTTP/2 client SSL_read failed. session_id=3 ssl_error=1",
            "[ERROR] SSL_accept() from client failed",
            "[ERROR] CONNECT tunnel recv() from client failed. session_id=7 error=10053",
        ],
    )
    add_body(
        doc,
        "이 로그들은 모두 동일한 의미가 아니다. SSL_accept 실패는 브라우저가 CONNECT 후 TLS 핸드셰이크를 완료하기 전에 연결을 닫은 경우가 많고, HTTP/2의 unexpected EOF와 Winsock 10053/10054는 프록시 재시작 또는 브라우저 연결 교체 과정에서 발생할 수 있다. 수정 후에는 이러한 정상 재연결 상황을 DEBUG로 낮추고, 실제 인증서·프로토콜 오류만 ERROR와 OpenSSL 상세 스택으로 남기도록 변경했다."
    )

    add_heading(doc, "5. 실제 업로드 검증 결과", level=1)
    add_body(
        doc,
        "Chrome 기반 ChatGPT 웹페이지에서 DOCX를 직접 업로드하여 메타데이터 요청, HTTP/2 raw PUT 스트림, DOCX 분석, upstream 전달까지 확인했다. 한글 파일명은 relay_runtime.log를 UTF-8로 읽을 때 정상 보존되었으며, 콘솔 깨짐은 출력 인코딩 문제로 분리해 수정했다."
    )
    add_matrix(
        doc,
        ["검증 항목", "결과"],
        [
            ("협상 프로토콜", "ALPN h2, AI TLS READY 확인"),
            ("원본 파일명", "5주차_TLS_MITM_구현사항_반영_정리.docx"),
            ("업로드 크기", "46,494 bytes (메타데이터와 raw PUT 일치)"),
            ("파일 형식", "DOCX / ZIP 컨테이너"),
            ("내부 엔트리", "18개"),
            ("추출 텍스트 바이트", "14,163 bytes"),
            ("SHA-256", "85b0a01e005558dba92a8b1ce0f33257175a524e789b790c55756922eac4d421"),
            ("정책 결과", "ALLOW"),
            ("원격 서버 결과", "HTTP 201 / UPLOAD FORWARDED"),
        ],
        [2500, 6860],
        font_size=9.3,
    )
    add_callout(
        doc,
        "검증 결론",
        "HTTPS 파일 업로드의 원본 파일명과 실제 파일 바이트를 복호화 구간에서 연결하고, 파일 구조를 검사한 뒤 안전한 경우에만 원격 AI 서비스로 전달하는 핵심 PoC 흐름이 동작했다.",
        fill="EAF4EE",
        accent=GREEN,
    )

    add_heading(doc, "6. 주요 변경 파일", level=1)
    add_matrix(
        doc,
        ["파일", "주요 변경", "역할"],
        [
            ("relay_proxy.c", "CONNECT 정책 분기, raw tunnel/MITM/Block 처리", "세션 진입점"),
            ("tls_mitm_engine.c", "동적 인증서, 이중 TLS, SNI, ALPN, HTTP/1.1 분석", "TLS 및 HTTP/1.1 핵심"),
            ("http2_engine.c", "HPACK, stream_id 상태, DATA 누적, 스트림 차단", "HTTP/2 분석"),
            ("http_parser.c", "Method/Path/Header/Body 추출", "HTTP 요청 평문화"),
            ("request_buffer.c / response_buffer.c", "완전한 요청·응답 경계 계산", "메시지 조립"),
            ("dlp_engine.c", "파일 업로드 여부 판단, 일반 채팅 제외", "검사 대상 축소"),
            ("upload_tracker.c", "메타데이터와 raw PUT 연결, UTF-8 파일명 복원", "업로드 상관관계"),
            ("file_analyzer.c", "DOCX/PDF/ZIP 검사, SHA-256, 구조·텍스트 바이트 추출", "파일 안전성 분석"),
            ("logger.c", "INFO 압축, 세션 식별, UTF-16 콘솔 출력", "운영 로그 가독성"),
            ("tls_intercept_policy.txt", "브라우저+AI 사이트만 MITM, DEFAULT IGNORE", "분석 범위 통제"),
        ],
        [2100, 4300, 2960],
        font_size=8.8,
    )

    add_heading(doc, "7. 남은 과제 및 향후 업데이트 계획", level=1)
    add_body(
        doc,
        "현재 PoC는 HTTPS 트래픽을 복호화하고 파일 업로드를 식별·검사·전달하는 핵심 경로를 확보했다. 프로젝트를 업무 환경에 적용 가능한 수준으로 완성하려면 다음 단계가 필요하다."
    )
    add_matrix(
        doc,
        ["우선순위", "업데이트 항목", "완료 기준"],
        [
            ("P0", "한글 콘솔 출력 수정본을 표준 Release에 반영하고 Chrome/Edge 재검증", "한글 파일명, h2 연결, 업로드 성공 로그가 동일 실행에서 정상"),
            ("P0", "raw tunnel의 10053/10054 정상 종료 로그도 DEBUG로 정리", "비 AI 사이트 이용 중 불필요한 ERROR 없음"),
            ("P1", "DOCX 본문을 실제 UTF-8 텍스트로 추출", "document.xml의 문단·표·머리글·각주를 보존한 검사 입력 생성"),
            ("P1", "PDF/ZIP 추출기 고도화", "Unicode PDF, 압축 파일 인코딩, 중첩 아카이브 제한 처리"),
            ("P1", "정책 엔진과 파일 분석 결과 연결", "형식·내용·크기 규칙에 따라 ALLOW/BLOCK 근거를 일관되게 기록"),
            ("P2", "보안 저장·보존 정책", "원문/추출문 ACL, 암호화, 보존 기간, 자동 삭제, 감사 추적"),
            ("P2", "성능·안정성", "대용량 스트리밍, backpressure, 동시 스트림 제한, timeout, 메모리 상한"),
            ("P2", "배포 운영", "Root CA 배포·회수·회전, 브라우저별 회귀 테스트, 정책 중앙 관리"),
            ("P3", "HTTP/3/QUIC 대응", "UDP/443 우회 정책 또는 HTTP/2 fallback 전략 결정"),
        ],
        [1000, 4430, 3930],
        font_size=8.7,
    )
    add_callout(
        doc,
        "권장 다음 스프린트",
        "우선 DOCX 실제 본문 추출과 차단 정책 연결을 완성하고, 동시에 raw tunnel 정상 종료 로그를 정리한다. 이후 PDF·ZIP 고도화와 운영 보안 통제를 단계적으로 추가한다.",
        fill="FFF8E8",
        accent=GOLD,
    )

    add_heading(doc, "8. 발표 요약", level=1)
    for item in [
        "기존 프록시는 HTTPS에서 암호화 바이트만 전달했지만, 이번 주에는 정책 기반 TLS MITM과 이중 TLS 세션을 구현했다.",
        "CONNECT 단계에서는 브라우저 프로세스와 AI 호스트를 기준으로 MITM 여부를 결정해 일반 웹 사용에는 영향을 주지 않도록 했다.",
        "복호화 이후 HTTP/1.1과 HTTP/2를 ALPN 결과에 따라 분기하고, URL·Header·Body를 요청 또는 stream_id 단위로 추출한다.",
        "ChatGPT 파일 업로드에서는 메타데이터와 raw PUT을 연결해 원본 한글 파일명, 파일 크기, DOCX 구조, SHA-256, 전송 상태를 하나의 upload_id로 검증했다.",
        "다음 단계는 실제 문서 본문 추출, 내용 기반 정책 차단, PDF·ZIP 고도화, 보안 저장 및 운영 안정성 확보다.",
    ]:
        add_bullet(doc, item, bullet_num)

    add_heading(doc, "결론", level=1)
    add_body(
        doc,
        "이번 주 구현으로 HTTPS 파일 업로드 트래픽을 ‘연결 수립 → TLS 복호화 → HTTP 평문 복원 → 파일 식별·분석 → 안전 파일 재전송’까지 추적할 수 있게 되었다. 특히 실제 ChatGPT HTTP/2 업로드에서 한글 원본 파일명과 raw PUT을 연결하고 DOCX 검사 결과 및 upstream 201 응답을 확인함으로써 핵심 아키텍처를 검증했다. 남은 작업은 분석 정확도와 운영 안전성을 높이는 방향으로 집중하며, 일반 웹 접근을 방해하지 않는 파일 전용 DLP Proxy라는 목표를 유지한다."
    )

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_report()
