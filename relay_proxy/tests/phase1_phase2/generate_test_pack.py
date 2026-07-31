#!/usr/bin/env python3
"""Generate deterministic Local DLP phase 1/2 fixtures without network access."""

from __future__ import annotations

import argparse
import binascii
import csv
import hashlib
import io
import json
from pathlib import Path
import shutil
import struct
import zipfile
import zlib
from xml.sax.saxutils import escape

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from pypdf import PdfReader, PdfWriter
from reportlab.lib.pagesizes import A4
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas


KOREAN = "이 문서는 로컬 DLP 1차 2차 기능 검증용 문서입니다."
ENGLISH = "LOCAL DLP PHASE ONE TWO TEST DOCUMENT"
MARKER = "DLP_TEST_CONFIDENTIAL_MARKER_2026"
ALL_TEXT = f"{KOREAN}\n{ENGLISH}\n{MARKER}"
ALLOW_KOREAN = "이 문서는 공개 테스트 자료이며 외부 공유가 허용됩니다."
ALLOW_ENGLISH = "LOCAL DLP PUBLIC ALLOW TEST DOCUMENT"
ALLOW_MARKER = "DLP_TEST_PUBLIC_MARKER_2026"
ALLOW_TEXT = f"{ALLOW_KOREAN}\n{ALLOW_ENGLISH}\n{ALLOW_MARKER}"
TEST_ARCHIVE_PASSWORD = "LocalDlp-Test-2026!"
SYNTHETIC_RRN_VALID = "000101-3000008"
SYNTHETIC_RRN_INVALID = "000101-3000009"
SYNTHETIC_CARD_VALID = "4111 1111 1111 1111"
# Every 13-16 digit suffix also fails Luhn; this prevents a shorter suffix from
# accidentally matching the detector while remaining an obvious dummy value.
SYNTHETIC_CARD_INVALID = "9999 9999 9999 9999"
SYNTHETIC_EMAIL = "dlp.synthetic@example.invalid"
SYNTHETIC_PHONE = "010-0000-0000"


def font_path() -> Path:
    candidates = [
        Path("C:/Windows/Fonts/malgun.ttf"),
        Path("C:/Windows/Fonts/gulim.ttc"),
        Path("/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc"),
        Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
    ]
    for candidate in candidates:
        if candidate.is_file():
            return candidate
    raise RuntimeError("No Korean-capable font found. Set up Malgun Gothic or Noto Sans CJK.")


def write_plain_files(output: Path) -> None:
    # These two fixtures are the policy-control pair used by the native
    # relay_proxy.exe integration test.  Keep the ALLOW fixture free of every
    # configured blocking keyword, including words used in explanatory prose.
    (output / "allow_public.txt").write_text(ALLOW_TEXT + "\n", encoding="utf-8-sig")
    (output / "block_confidential.txt").write_text(ALL_TEXT + "\n", encoding="utf-8-sig")
    # Encoding-regression fixtures for the native text normalizer.  CP949 and
    # BOM-less UTF-16LE are deliberately written as raw bytes so an accidental
    # UTF-8-only implementation cannot make these tests pass.
    (output / "block_confidential_cp949.txt").write_bytes((ALL_TEXT + "\n").encode("cp949"))
    (output / "block_confidential_utf8_nobom.txt").write_bytes((ALL_TEXT + "\n").encode("utf-8"))
    utf16le = (ALL_TEXT + "\n").encode("utf-16le")
    (output / "block_confidential_utf16le_nobom.txt").write_bytes(utf16le)
    (output / "block_confidential_utf16le_bom.txt").write_bytes(b"\xff\xfe" + utf16le)
    (output / "block_confidential_utf16be_bom.txt").write_bytes(
        b"\xfe\xff" + (ALL_TEXT + "\n").encode("utf-16be")
    )
    (output / "block_confidential_utf32le_bom.txt").write_bytes(
        b"\xff\xfe\x00\x00" + (ALL_TEXT + "\n").encode("utf-32le")
    )
    (output / "block_confidential_utf32be_bom.txt").write_bytes(
        b"\x00\x00\xfe\xff" + (ALL_TEXT + "\n").encode("utf-32be")
    )

    # The remaining fixtures exercise format-specific extraction.  They are
    # intentionally policy-sensitive because MARKER contains CONFIDENTIAL.
    (output / "safe_sample.txt").write_text(ALL_TEXT + "\n", encoding="utf-8-sig")
    (output / "safe_sample.json").write_text(
        json.dumps({"korean": KOREAN, "english": ENGLISH, "marker": MARKER}, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    (output / "safe_sample.xml").write_text(
        f'<?xml version="1.0" encoding="UTF-8"?>\n<dlp-test><korean>{escape(KOREAN)}</korean><english>{ENGLISH}</english><marker>{MARKER}</marker></dlp-test>\n',
        encoding="utf-8",
    )
    (output / "safe_sample.csv").write_text(
        "field,value\n" + f'korean,"{KOREAN}"\nenglish,"{ENGLISH}"\nmarker,"{MARKER}"\n',
        encoding="utf-8-sig",
    )
    (output / "safe_sample.py").write_text(
        '# Synthetic DLP source-code fixture\n' + f'DLP_TEST_TEXT = "{ENGLISH}"\nDLP_TEST_MARKER = "{MARKER}"\n',
        encoding="utf-8",
    )

    # Policy-engine controls use conspicuously synthetic values.  The resident
    # number passes only the implemented date/checksum validator and the card is
    # the ubiquitous Luhn test value; neither fixture represents a real person
    # or an active payment instrument.
    policy_files = {
        "policy_block_rule100.txt": "SYNTHETIC POLICY KEYWORD TEST\n표식: 대외비\n",
        "policy_block_rule101.txt": "SYNTHETIC POLICY KEYWORD TEST\n표식: 기밀\n",
        "policy_block_rule110_rrn.txt": (
            "SYNTHETIC CHECKSUM TEST VALUE - NOT A REAL PERSON\n" + SYNTHETIC_RRN_VALID + "\n"
        ),
        "policy_block_rule111_card.txt": (
            "SYNTHETIC LUHN TEST VALUE - NOT AN ACTIVE PAYMENT CARD\n" + SYNTHETIC_CARD_VALID + "\n"
        ),
        "policy_allow_invalid_rrn.txt": (
            "SYNTHETIC INVALID RESIDENT CHECKSUM CONTROL\n" + SYNTHETIC_RRN_INVALID + "\n"
        ),
        "policy_allow_invalid_card.txt": (
            "SYNTHETIC INVALID LUHN CONTROL\n" + SYNTHETIC_CARD_INVALID + "\n"
        ),
        "policy_log_only_email.txt": "SYNTHETIC CONTACT TEST\n" + SYNTHETIC_EMAIL + "\n",
        "policy_log_only_phone.txt": "SYNTHETIC CONTACT TEST\n" + SYNTHETIC_PHONE + "\n",
    }
    for filename, content in policy_files.items():
        (output / filename).write_text(content, encoding="utf-8")


def write_docx(output: Path) -> None:
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.85)
    section.bottom_margin = Inches(0.85)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)
    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Malgun Gothic"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
    normal.font.size = Pt(11)

    heading = document.add_paragraph()
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = heading.add_run("Local DLP 문서 본문 추출 테스트")
    run.bold = True
    run.font.size = Pt(20)
    run.font.name = "Malgun Gothic"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
    document.add_paragraph(KOREAN)
    document.add_paragraph(ENGLISH)
    document.add_paragraph(MARKER)
    table = document.add_table(rows=4, cols=2)
    table.style = "Light Shading Accent 1"
    values = [("Field", "Value"), ("Korean", KOREAN), ("English", ENGLISH), ("Marker", MARKER)]
    for row, values_row in zip(table.rows, values):
        for cell, value in zip(row.cells, values_row):
            cell.text = value
    document.core_properties.title = "Local DLP extraction fixture"
    document.core_properties.subject = "Synthetic test data"
    document.save(output / "safe_sample.docx")

    # Isolate the embedded-payload fail-closed rule from content policy.  The
    # visible document is public/benign; only the OPC embedded package is
    # unsupported.  A relationship and content type make the synthetic part a
    # standards-shaped Office package payload instead of an arbitrary ZIP file.
    public_document = Document()
    public_document.add_paragraph(ALLOW_TEXT)
    temporary = output / ".embedded_opaque_base.docx"
    public_document.save(temporary)
    target = output / "embedded_opaque_payload.docx"
    with zipfile.ZipFile(temporary, "r") as source, zipfile.ZipFile(
        target, "w", compression=zipfile.ZIP_DEFLATED
    ) as destination:
        for info in source.infolist():
            data = source.read(info.filename)
            if info.filename == "[Content_Types].xml":
                data = data.replace(
                    b"</Types>",
                    b'<Default Extension="bin" ContentType="application/octet-stream"/></Types>',
                )
            elif info.filename == "word/_rels/document.xml.rels":
                data = data.replace(
                    b"</Relationships>",
                    b'<Relationship Id="rIdDlpOpaque" '
                    b'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/package" '
                    b'Target="embeddings/opaque.bin"/></Relationships>',
                )
            destination.writestr(info, data)
        destination.writestr(
            "word/embeddings/opaque.bin",
            b"LOCAL_DLP_SYNTHETIC_OPAQUE_EMBEDDED_PACKAGE\x00\x01\x02\xff",
        )
    temporary.unlink()


def write_images(output: Path) -> Image.Image:
    width, height = 1800, 760
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    font = ImageFont.truetype(str(font_path()), 68)
    small = ImageFont.truetype(str(font_path()), 54)
    draw.rectangle((38, 38, width - 38, height - 38), outline="#2563EB", width=8)
    draw.text((92, 112), "Local DLP OCR TEST", font=font, fill="black")
    draw.text((92, 270), KOREAN, font=small, fill="black")
    draw.text((92, 410), ENGLISH, font=small, fill="black")
    draw.text((92, 550), MARKER, font=small, fill="black")
    image.save(output / "safe_scan.png", optimize=True)
    image.save(output / "safe_scan.jpg", quality=96, subsampling=0)
    image.save(output / "safe_scan.gif")
    image.save(output / "safe_scan.bmp")
    image.save(output / "safe_scan.tiff", compression="tiff_deflate")

    # The blocking marker exists only on frame/page 2.  These fixtures catch
    # regressions where Windows OCR silently inspects only the first frame.
    benign = Image.new("RGB", (width, height), "white")
    benign_draw = ImageDraw.Draw(benign)
    benign_draw.rectangle((38, 38, width - 38, height - 38), outline="#16A34A", width=8)
    benign_draw.text((92, 170), "LOCAL DLP FRAME ONE", font=font, fill="black")
    benign_draw.text((92, 340), ALLOW_ENGLISH, font=small, fill="black")
    benign_draw.text((92, 500), ALLOW_MARKER, font=small, fill="black")
    benign.save(
        output / "marker_second_frame.gif",
        save_all=True,
        append_images=[image],
        duration=[500, 500],
        loop=0,
        disposal=2,
    )
    benign.save(
        output / "marker_second_page.tiff",
        save_all=True,
        append_images=[image],
        compression="tiff_deflate",
    )
    return image


def write_pdfs(output: Path, scan_image: Image.Image) -> None:
    target = output / "safe_text.pdf"
    korean_font = font_path()
    pdfmetrics.registerFont(TTFont("DlpKorean", str(korean_font)))
    pdf = canvas.Canvas(str(target), pagesize=A4)
    pdf.setTitle("Local DLP text PDF fixture")
    pdf.setFont("DlpKorean", 18)
    pdf.drawString(56, 780, "Local DLP PDF 본문 추출 테스트")
    pdf.setFont("DlpKorean", 13)
    pdf.drawString(56, 720, KOREAN)
    pdf.drawString(56, 680, ENGLISH)
    pdf.drawString(56, 640, MARKER)
    pdf.save()

    encrypted_target = output / "encrypted_confidential.pdf"
    reader = PdfReader(str(target))
    writer = PdfWriter()
    for page in reader.pages:
        writer.add_page(page)
    writer.add_metadata({"/Title": "Synthetic encrypted Local DLP fixture"})
    writer.encrypt(TEST_ARCHIVE_PASSWORD)
    with encrypted_target.open("wb") as stream:
        writer.write(stream)

    scan_target = output / "safe_scan.pdf"
    temporary_png = output / ".safe_scan_pdf_source.png"
    scan_image.save(temporary_png)
    scan_pdf = canvas.Canvas(str(scan_target), pagesize=A4)
    scan_pdf.setTitle("Local DLP scanned PDF OCR fixture")
    scan_pdf.drawImage(str(temporary_png), 28, 265, width=540, height=228, preserveAspectRatio=True, mask="auto")
    scan_pdf.save()
    temporary_png.unlink()


def write_hwpx(output: Path) -> None:
    target = output / "safe_sample.hwpx"
    section_xml = f'''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<hs:sec xmlns:hs="http://www.hancom.co.kr/hwpml/2011/section" xmlns:hp="http://www.hancom.co.kr/hwpml/2011/paragraph">
  <hp:p id="0" paraPrIDRef="0" styleIDRef="0"><hp:run charPrIDRef="0"><hp:t>{escape(KOREAN)}</hp:t></hp:run></hp:p>
  <hp:p id="1" paraPrIDRef="0" styleIDRef="0"><hp:run charPrIDRef="0"><hp:t>{ENGLISH}</hp:t></hp:run></hp:p>
  <hp:p id="2" paraPrIDRef="0" styleIDRef="0"><hp:run charPrIDRef="0"><hp:t>{MARKER}</hp:t></hp:run></hp:p>
</hs:sec>'''
    content_hpf = '''<?xml version="1.0" encoding="UTF-8"?>
<opf:package xmlns:opf="http://www.idpf.org/2007/opf" version="3.0" unique-identifier="bookid">
  <opf:metadata><opf:title>Local DLP HWPX fixture</opf:title></opf:metadata>
  <opf:manifest><opf:item id="section0" href="Contents/section0.xml" media-type="application/xml"/></opf:manifest>
  <opf:spine><opf:itemref idref="section0"/></opf:spine>
</opf:package>'''
    manifest = '''<?xml version="1.0" encoding="UTF-8"?>
<manifest xmlns="urn:oasis:names:tc:opendocument:xmlns:manifest:1.0">
  <file-entry media-type="application/hwp+zip" full-path="/"/>
  <file-entry media-type="application/xml" full-path="Contents/content.hpf"/>
  <file-entry media-type="application/xml" full-path="Contents/section0.xml"/>
</manifest>'''
    with zipfile.ZipFile(target, "w") as archive:
        archive.writestr("mimetype", "application/hwp+zip", compress_type=zipfile.ZIP_STORED)
        archive.writestr("META-INF/manifest.xml", manifest, compress_type=zipfile.ZIP_DEFLATED)
        archive.writestr("Contents/content.hpf", content_hpf, compress_type=zipfile.ZIP_DEFLATED)
        archive.writestr("Contents/section0.xml", section_xml, compress_type=zipfile.ZIP_DEFLATED)


def _cfb_directory_entry(
    name: str,
    entry_type: int,
    start_sector: int,
    stream_size: int,
    child: int = 0xFFFFFFFF,
    right: int = 0xFFFFFFFF,
) -> bytes:
    entry = bytearray(128)
    encoded_name = (name + "\0").encode("utf-16le")
    if len(encoded_name) > 64:
        raise ValueError(f"CFB directory name is too long: {name}")
    entry[: len(encoded_name)] = encoded_name
    struct.pack_into("<H", entry, 64, len(encoded_name))
    entry[66] = entry_type
    entry[67] = 1  # black node
    struct.pack_into("<III", entry, 68, 0xFFFFFFFF, right, child)
    struct.pack_into("<I", entry, 116, start_sector)
    struct.pack_into("<Q", entry, 120, stream_size)
    return bytes(entry)


def _write_minimal_cfb(path: Path, streams: list[tuple[str, bytes]]) -> None:
    """Write a deterministic CFB v3 file using regular 512-byte sectors.

    Stream sizes remain below the normal mini-stream cutoff.  A root mini stream
    is intentionally absent; readers, including the production analyzer, then
    follow the regular FAT chains.  This keeps the fixture compact while
    retaining a standards-shaped compound-file header and directory.
    """

    sector_size = 512
    free = 0xFFFFFFFF
    end = 0xFFFFFFFE
    fat_sector = 0xFFFFFFFD
    sectors: list[bytes] = []
    directory_sector_id = 0
    sectors.append(bytes(sector_size))
    stream_layout: list[tuple[str, int, int, bytes]] = []

    for name, content in streams:
        start = len(sectors)
        count = max(1, (len(content) + sector_size - 1) // sector_size)
        padded = content + bytes(count * sector_size - len(content))
        for index in range(count):
            sectors.append(padded[index * sector_size : (index + 1) * sector_size])
        stream_layout.append((name, start, len(content), padded))

    fat_sector_id = len(sectors)
    if fat_sector_id >= 128:
        raise RuntimeError("Minimal CFB fixture exceeded one FAT sector")

    directory = bytearray(sector_size)
    directory[0:128] = _cfb_directory_entry(
        "Root Entry", 5, end, 0, child=1 if stream_layout else free
    )
    for index, (name, start, size, _padded) in enumerate(stream_layout, start=1):
        right = index + 1 if index < len(stream_layout) else free
        directory[index * 128 : (index + 1) * 128] = _cfb_directory_entry(
            name, 2, start, size, right=right
        )
    sectors[directory_sector_id] = bytes(directory)

    fat = [free] * 128
    fat[directory_sector_id] = end
    for _name, start, size, _padded in stream_layout:
        count = max(1, (size + sector_size - 1) // sector_size)
        for offset in range(count):
            sector_id = start + offset
            fat[sector_id] = end if offset + 1 == count else sector_id + 1
    fat[fat_sector_id] = fat_sector
    sectors.append(struct.pack("<128I", *fat))

    header = bytearray(sector_size)
    header[:8] = bytes.fromhex("D0CF11E0A1B11AE1")
    struct.pack_into("<HHHHH", header, 24, 0x003E, 3, 0xFFFE, 9, 6)
    struct.pack_into("<I", header, 40, 0)  # directory sectors for CFB v3
    struct.pack_into("<I", header, 44, 1)  # FAT sector count
    struct.pack_into("<I", header, 48, directory_sector_id)
    struct.pack_into("<I", header, 56, 4096)
    struct.pack_into("<I", header, 60, end)
    struct.pack_into("<I", header, 64, 0)
    struct.pack_into("<I", header, 68, end)
    struct.pack_into("<I", header, 72, 0)
    for index in range(109):
        struct.pack_into("<I", header, 76 + index * 4, free)
    struct.pack_into("<I", header, 76, fat_sector_id)
    path.write_bytes(bytes(header) + b"".join(sectors))


def write_minimal_hwp(output: Path) -> None:
    """Create a minimal uncompressed HWP 5.x compound document.

    The fixture contains the standard FileHeader signature and one Section0
    HWPTAG_PARA_TEXT record.  It is synthetic and does not depend on Hancom
    Office, so it must never be presented as a vendor round-trip compatibility
    corpus.
    """

    file_header = bytearray(256)
    signature = b"HWP Document File"
    file_header[: len(signature)] = signature
    struct.pack_into("<I", file_header, 32, 0x05000302)
    struct.pack_into("<I", file_header, 36, 0)  # uncompressed, not encrypted

    paragraph = ALL_TEXT.encode("utf-16le")
    if len(paragraph) >= 0xFFF:
        raise RuntimeError("HWP paragraph fixture unexpectedly requires an extended record header")
    record_header = 0x43 | (len(paragraph) << 20)  # HWPTAG_PARA_TEXT, level 0
    section = struct.pack("<I", record_header) + paragraph
    _write_minimal_cfb(
        output / "minimal_hwp5_confidential.hwp",
        [("FileHeader", bytes(file_header)), ("Section0", section)],
    )


def write_minimal_legacy_office(output: Path) -> None:
    """Create bounded CFB parser fixtures for the legacy Office fail-closed path.

    These are intentionally not represented as Microsoft Office round-trip
    documents.  They exercise CFB recognition, expected stream naming, bounded
    UTF-16 evidence extraction, and the production rule that legacy formats
    remain blocked after partial inspection.
    """

    encoded = (ALL_TEXT + "\n").encode("utf-16le")
    # CFB streams below 4096 bytes normally live in the mini stream.  Padding
    # these legacy evidence streams to the cutoff keeps them on the regular FAT
    # used by this compact writer and makes the compound-file structure valid
    # for independent CFB readers as well as the production parser.
    payload = encoded + bytes(4096 - len(encoded))
    fixtures = (
        ("synthetic_legacy_confidential.doc", "WordDocument"),
        ("synthetic_legacy_confidential.xls", "Workbook"),
        ("synthetic_legacy_confidential.ppt", "PowerPoint Document"),
    )
    for filename, stream_name in fixtures:
        _write_minimal_cfb(output / filename, [(stream_name, payload)])


def _zipcrypto_crc_update(value: int, byte: int) -> int:
    crc = value ^ byte
    for _ in range(8):
        crc = (crc >> 1) ^ (0xEDB88320 if crc & 1 else 0)
    return crc & 0xFFFFFFFF


def _zipcrypto_encrypt(data: bytes, password: bytes) -> bytes:
    keys = [0x12345678, 0x23456789, 0x34567890]

    def update(byte: int) -> None:
        keys[0] = _zipcrypto_crc_update(keys[0], byte)
        keys[1] = ((keys[1] + (keys[0] & 0xFF)) * 134775813 + 1) & 0xFFFFFFFF
        keys[2] = _zipcrypto_crc_update(keys[2], (keys[1] >> 24) & 0xFF)

    for byte in password:
        update(byte)
    encrypted = bytearray()
    for byte in data:
        temporary = (keys[2] | 2) & 0xFFFFFFFF
        mask = ((temporary * (temporary ^ 1)) >> 8) & 0xFF
        encrypted.append(byte ^ mask)
        update(byte)
    return bytes(encrypted)


def write_encrypted_zip(output: Path) -> None:
    """Write a deterministic, genuine traditional ZipCrypto archive."""

    target = output / "encrypted_confidential.zip"
    filename = b"confidential.txt"
    payload = (ALL_TEXT + "\n").encode("utf-8")
    password = TEST_ARCHIVE_PASSWORD.encode("utf-8")
    crc = binascii.crc32(payload) & 0xFFFFFFFF
    compressor = zlib.compressobj(level=9, wbits=-15)
    compressed = compressor.compress(payload) + compressor.flush()
    deterministic = hashlib.sha256(filename + password + payload).digest()[:11]
    encryption_header = deterministic + bytes([(crc >> 24) & 0xFF])
    encrypted_payload = _zipcrypto_encrypt(encryption_header + compressed, password)

    flags = 0x0001
    method = 8
    dos_time = (12 << 11)  # 12:00:00
    dos_date = ((2026 - 1980) << 9) | (1 << 5) | 1
    compressed_size = len(encrypted_payload)
    local = struct.pack(
        "<IHHHHHIIIHH",
        0x04034B50, 20, flags, method, dos_time, dos_date,
        crc, compressed_size, len(payload), len(filename), 0,
    ) + filename + encrypted_payload
    central = struct.pack(
        "<IHHHHHHIIIHHHHHII",
        0x02014B50, 20, 20, flags, method, dos_time, dos_date,
        crc, compressed_size, len(payload), len(filename), 0, 0, 0, 0, 0, 0,
    ) + filename
    end = struct.pack(
        "<IHHHHIIH",
        0x06054B50, 0, 0, 1, 1, len(central), len(local), 0,
    )
    target.write_bytes(local + central + end)

    with zipfile.ZipFile(target) as archive:
        recovered = archive.read(filename.decode("ascii"), pwd=password)
    if recovered != payload:
        raise RuntimeError("Encrypted ZIP fixture failed password round-trip verification")


def write_zip(output: Path) -> None:
    with zipfile.ZipFile(output / "safe_bundle.zip", "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for filename in ("safe_sample.txt", "safe_sample.json", "safe_sample.docx"):
            archive.write(output / filename, arcname=f"documents/{filename}")


def write_manifest(output: Path) -> None:
    fixtures = [
        ("allow_public.txt", "text/plain", "text", ALLOW_MARKER),
        ("block_confidential.txt", "text/plain", "text", MARKER),
        ("block_confidential_cp949.txt", "text/plain", "text", MARKER),
        ("block_confidential_utf8_nobom.txt", "text/plain", "text", MARKER),
        ("block_confidential_utf16le_nobom.txt", "text/plain", "text", MARKER),
        ("block_confidential_utf16le_bom.txt", "text/plain", "text", MARKER),
        ("block_confidential_utf16be_bom.txt", "text/plain", "text", MARKER),
        ("block_confidential_utf32le_bom.txt", "text/plain", "text", MARKER),
        ("block_confidential_utf32be_bom.txt", "text/plain", "text", MARKER),
        ("safe_sample.txt", "text/plain", "text", MARKER),
        ("safe_sample.csv", "text/csv", "text", MARKER),
        ("safe_sample.json", "application/json", "text", MARKER),
        ("safe_sample.xml", "application/xml", "text", MARKER),
        ("safe_sample.py", "text/x-python", "text", MARKER),
        ("policy_block_rule100.txt", "text/plain", "text", "대외비"),
        ("policy_block_rule101.txt", "text/plain", "text", "기밀"),
        ("policy_block_rule110_rrn.txt", "text/plain", "text", SYNTHETIC_RRN_VALID),
        ("policy_block_rule111_card.txt", "text/plain", "text", SYNTHETIC_CARD_VALID),
        ("policy_allow_invalid_rrn.txt", "text/plain", "text", SYNTHETIC_RRN_INVALID),
        ("policy_allow_invalid_card.txt", "text/plain", "text", SYNTHETIC_CARD_INVALID),
        ("policy_log_only_email.txt", "text/plain", "text", SYNTHETIC_EMAIL),
        ("policy_log_only_phone.txt", "text/plain", "text", SYNTHETIC_PHONE),
        ("safe_sample.docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document", "document", MARKER),
        ("safe_sample.xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "document", MARKER),
        ("safe_sample.pptx", "application/vnd.openxmlformats-officedocument.presentationml.presentation", "document", MARKER),
        ("safe_text.pdf", "application/pdf", "document", MARKER),
        ("safe_scan.pdf", "application/pdf", "ocr", MARKER),
        ("safe_sample.hwpx", "application/hwp+zip", "document", MARKER),
        ("safe_bundle.zip", "application/zip", "archive", MARKER),
        ("safe_scan.png", "image/png", "ocr", MARKER),
        ("safe_scan.jpg", "image/jpeg", "ocr", MARKER),
        ("safe_scan.gif", "image/gif", "ocr", MARKER),
        ("safe_scan.bmp", "image/bmp", "ocr", MARKER),
        ("safe_scan.tiff", "image/tiff", "ocr", MARKER),
    ]
    manifest = {
        "schema": 2,
        "synthetic_test_data": True,
        "expected": {
            "korean": KOREAN,
            "english": ENGLISH,
            "marker": MARKER,
            "allow_marker": ALLOW_MARKER,
            "policy_keyword_100": "대외비",
            "policy_keyword_101": "기밀",
        },
        "fixtures": [
            {"file": filename, "mime": mime, "analysis": analysis, "must_contain": [required_marker]}
            for filename, mime, analysis, required_marker in fixtures
        ],
        "native_only_fixtures": [
            {
                "file": "marker_second_frame.gif",
                "mime": "image/gif",
                "expected_format": "GIF",
                "expected_action": "BLOCK",
                "expected_rule": 102,
                "must_contain": [MARKER],
                "marker_frame": 2,
                "provenance": "two-frame animated GIF; blocking marker appears only on frame 2",
            },
            {
                "file": "marker_second_page.tiff",
                "mime": "image/tiff",
                "expected_format": "TIFF",
                "expected_action": "BLOCK",
                "expected_rule": 102,
                "must_contain": [MARKER],
                "marker_frame": 2,
                "provenance": "two-page TIFF; blocking marker appears only on page 2",
            },
            {
                "file": "embedded_opaque_payload.docx",
                "mime": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                "expected_format": "DOCX/DOCM",
                "expected_action": "BLOCK",
                "expected_reason": "embedded package payload cannot be fully inspected",
                "provenance": "benign DOCX with a related word/embeddings/opaque.bin package part",
            },
            {
                "file": "encrypted_confidential.pdf",
                "mime": "application/pdf",
                "expected_format": "PDF",
                "expected_action": "BLOCK",
                "expected_reason": "encrypted PDF",
                "password": TEST_ARCHIVE_PASSWORD,
                "provenance": "genuine pypdf password-encrypted PDF; password round-trip verified",
            },
            {
                "file": "synthetic_legacy_confidential.doc",
                "mime": "application/msword",
                "expected_format": "DOC",
                "expected_action": "BLOCK",
                "expected_reason": "legacy OLE text evidence extracted partially",
                "must_contain": [MARKER],
                "provenance": "synthetic minimal CFB with a WordDocument stream; not a Word vendor round-trip file",
            },
            {
                "file": "synthetic_legacy_confidential.xls",
                "mime": "application/vnd.ms-excel",
                "expected_format": "XLS",
                "expected_action": "BLOCK",
                "expected_reason": "legacy OLE text evidence extracted partially",
                "must_contain": [MARKER],
                "provenance": "synthetic minimal CFB with a Workbook stream; not an Excel vendor round-trip file",
            },
            {
                "file": "synthetic_legacy_confidential.ppt",
                "mime": "application/vnd.ms-powerpoint",
                "expected_format": "PPT",
                "expected_action": "BLOCK",
                "expected_reason": "legacy OLE text evidence extracted partially",
                "must_contain": [MARKER],
                "provenance": "synthetic minimal CFB with a PowerPoint Document stream; not a PowerPoint vendor round-trip file",
            },
            {
                "file": "minimal_hwp5_confidential.hwp",
                "mime": "application/x-hwp",
                "expected_format": "HWP",
                "expected_action": "BLOCK",
                "expected_rule": 102,
                "must_contain": [MARKER],
                "provenance": "synthetic minimal HWP 5.x CFB with FileHeader and Section0 PARA_TEXT",
            },
            {
                "file": "encrypted_confidential.zip",
                "mime": "application/zip",
                "expected_format": "ZIP",
                "expected_action": "BLOCK",
                "expected_reason": "encrypted archive entry",
                "password": TEST_ARCHIVE_PASSWORD,
                "provenance": "deterministic traditional ZipCrypto archive; password round-trip verified",
            },
        ],
        "not_generated": [
            {
                "formats": ["doc", "xls", "ppt"],
                "reason": "Vendor round-trip generation is optional; the Windows Office helper records installed/running/watchdog status.",
            },
            {
                "formats": ["encrypted office"],
                "reason": "Password-protected Office generation is optional and requires a closed, responsive Word COM instance.",
            },
        ],
    }
    (output.parent / "manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", type=Path, default=Path(__file__).with_name("fixtures"))
    args = parser.parse_args()
    output = args.output.resolve()
    if output.name.casefold() != "fixtures" or "ai_dlp" not in output.parent.name.casefold():
        raise RuntimeError(
            "Refusing to clean an unsafe output directory. Use an AI_DLP-named pack directory with a fixtures child."
        )
    if output.exists():
        for child in output.iterdir():
            if child.is_file():
                child.unlink()
            elif child.is_dir():
                shutil.rmtree(child)
    output.mkdir(parents=True, exist_ok=True)
    write_plain_files(output)
    write_docx(output)
    write_zip(output)
    image = write_images(output)
    write_pdfs(output, image)
    write_hwpx(output)
    write_minimal_hwp(output)
    write_minimal_legacy_office(output)
    write_encrypted_zip(output)
    write_manifest(output)
    print(json.dumps({"output": str(output), "status": "base fixtures generated"}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
